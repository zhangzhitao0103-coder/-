"""生成主视觉清单：每校文件名、建议画面、AI 生图提示词。

图片放到 templates/assets/photos/<文件名>，再运行 python3 scripts/apply_heroes.py 即可换图出图。
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from order import ORDER

ROOT = Path(__file__).resolve().parent.parent
notes = dict(re.findall(r'"(\w+)": \(.*?\),\s*#\s*(.+)', (ROOT / "scripts" / "make_hero.py").read_text(encoding="utf-8")))
notes["bit"] = "中关村校区主楼"

wb = Workbook()
ws = wb.active
ws.title = "主视觉清单"
ws.append(["序号", "校名", "图片文件名", "建议画面", "AI 生图提示词", "状态"])
for i, sid in enumerate(ORDER, 1):
    d = json.loads((ROOT / "data" / f"{sid}.json").read_text(encoding="utf-8"))
    scene = notes.get(sid, "")
    if not any(k in scene for k in "楼门塔堂馆院湖山园路") or "现代" in scene:
        scene = d["landmarks"][0][0]  # 示意图注释太笼统时，用数据里的第一个校园风物
    prompt = (f"{d['name']}{scene}，{d['city'].replace('·', '')}，校园建筑正面广角，晴天蓝天白云，"
              f"清新明亮，写实摄影风格（或水彩插画风格），画面无文字、无人物特写，横幅构图，比例 2.4:1")
    has = any((ROOT / "templates" / "assets" / "photos" / f"{sid}{ext}").exists() for ext in (".jpg", ".jpeg", ".png", ".webp"))
    ws.append([i, d["name"], f"{sid}.jpg", scene, prompt, "已替换" if has else "待提供"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="0B6B3F")
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDEF", [6, 20, 14, 22, 70, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C2"

ws2 = wb.create_sheet("图片要求")
for r in [
    ["尺寸", "横图，至少 2000×840 像素（海报图框约 2.4:1，会居中裁切填满）"],
    ["格式", "jpg / png / webp 均可，文件名用「图片文件名」列，如 thu.jpg"],
    ["内容", "学校标志性建筑或校门，正面、光线好；避免人物特写和大段文字"],
    ["版权", "实景照片需有使用授权：自己拍摄、学校授权，或 CC 协议图片（注明作者）；不要直接用学校官网或新闻图片"],
    ["AI 插画", "可直接把「AI 生图提示词」列粘贴到生图工具；同一风格最好用同一工具和参数，保持系列统一"],
    ["放置位置", "college-posters/templates/assets/photos/"],
    ["换图命令", "python3 scripts/apply_heroes.py（只替换有图片的学校，其余保留示意插画）"],
]:
    ws2.append(r)
ws2.column_dimensions["A"].width = 10
ws2.column_dimensions["B"].width = 90
dst = ROOT / "data" / "主视觉清单.xlsx"
wb.save(dst)
print(dst, len(ORDER))
