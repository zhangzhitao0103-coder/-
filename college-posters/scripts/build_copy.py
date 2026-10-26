"""生成朋友圈文案：115 所 × 6 版，按发布顺序排期（2026-11-02 周一起，每周三为家长干货日）。
输出 data/朋友圈文案.xlsx。手写素材见 copy_hooks.py，其余字段全部取自 data/<id>.json（与海报一致）。"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from copy_hooks import CAT, H  # noqa: E402
from order import ORDER  # noqa: E402

START = dt.date(2026, 11, 2)
WEEK = "一二三四五六日"
TITLE = "🎓 每天认识一所大学｜{}"

AUD = {"工": "理工爱好者值得了解的大学", "综": "一所选择丰富的综合性大学", "师": "未来老师值得了解的大学",
       "农": "热爱自然的你值得了解", "医": "想学医的你值得了解", "财": "对经济感兴趣的你值得了解",
       "外": "语言爱好者值得了解", "法": "想学法律的你值得了解", "艺": "有艺术梦想的你值得了解",
       "体": "热爱运动的你值得了解", "军": "有从军梦想的你值得了解"}
LEARN = {"工": "🔬 这里重视从原理走向实践，用知识解决真问题",
         "综": "🔬 学科齐全、选择多，更看重扎实的基础",
         "师": "🔬 这里不只学知识，也学怎样把知识教明白",
         "农": "🔬 课堂连着田野和实验室，科学就用在土地上",
         "医": "🔬 学医路长，要有扎实的生化基础和一颗仁心",
         "财": "🔬 在这里，数学和逻辑是读懂经济的钥匙",
         "外": "🔬 语言是工具，更是理解世界的窗口",
         "法": "🔬 严密的逻辑和清晰的表达，同样重要",
         "艺": "🔬 专业路上，热爱与积累之外，也要文化课托底",
         "体": "🔬 体育之路，既要汗水，也要文化课托底",
         "军": "🔬 军校学习，专业训练与军事素养并重"}

ENC_A = ["今天多弄懂一个原理、多解决一道难题，都在为未来的选择积累底气",
         "把每一次作业当成一次小考认真对待，目标会越来越近",
         "名校不是遥远的风景，而是一道道题、一次次复盘铺出来的路",
         "今天的错题本，就是明天的提分本",
         "目标定得高一点，行动落得细一点，每一天都算数",
         "先把今天的课听懂、作业写对，就是向理想大学迈出的一步",
         "别急着和别人比，和昨天的自己比，进步就看得见",
         "补一补薄弱学科，拔一拔优势学科，分数会给你答案",
         "高中三年很长也很短，每一天的认真都不会白费",
         "把目标写下来贴在书桌前，累的时候抬头看一看"]
ENC_B = ["你羡慕的校园，也可以是你未来四年的日常",
         "今天的努力，是为了将来能走进这样的校园",
         "也许几年后，你就是这里的新生",
         "把这所大学存进心愿单，然后用行动靠近它",
         "风景很美，最美的是你走进去的那一天",
         "愿你有一天，能亲眼看看这里的风景",
         "好大学在远方，路就在今天的书桌上"]
ENC_C = ["认真完成一次推导，坚持弄懂一个问题，让每天的积累都有方向",
         "历史是一步步走出来的，成绩也是",
         "每一所名校都有来路，你的目标也可以从今天的小事开始",
         "先定一个小目标：今天弄懂一道错题",
         "把大目标拆成每天的小任务，一步一步，终会抵达",
         "前人能在艰难中坚持，我们也能在平凡日子里努力"]
ENC_Q = ["不妨从今天最感兴趣的一道题开始，让好奇心慢慢长成本领",
         "兴趣是最好的老师，把它带进每一堂课",
         "把喜欢的事研究透，就是最好的学习",
         "先把相关学科学扎实，兴趣才能变成真本领",
         "带着好奇心去学习，你会发现每一科都有意思"]
MOTTO_TAIL = ["提醒我们把基础学扎实", "值得抄在笔记本扉页上",
              "也可以是你的座右铭", "这就是做学问的态度"]
CTA = ["想知道孩子现在离{s}还差多少分？私信我，免费做一次学情分析",
       "孩子的选科和目标院校还没想清楚？私信我，一起做个规划",
       "想了解孩子现在的成绩能冲哪些好大学？私信我，免费帮你分析"]
CTA_MIL = "军校还要过政审和体检，想了解报考条件？私信我"


def motto(d):
    parts = d["motto"].split()
    return "".join(parts) if all(len(p) <= 2 for p in parts) else "，".join(parts)


def motto_label(d):
    lab = d.get("motto_label", "校训")
    return f"{d['short']}精神" if lab == "精神" else lab


def motto_note(d):
    n = re.sub(r"（.*?）", "", d.get("motto_note") or "")
    return "" if (not n or "寄语" in n or "待核" in n or len(n) > 18) else n


def ident(d):
    return " · ".join(t for t in ("985", "211", "双一流") if t in d["tags"])


def rank(d):
    if not d["rankings"]:
        return ""
    r = d["rankings"][0]
    if r["label"].startswith("软科"):
        return f"软科{r['year']}{r['label'][2:]}第{r['value']}"
    if r["label"] == "国家双一流建设学科" and "双一流" not in d["eval_note"]:
        return f"{r['value']}个国家双一流建设学科"
    return ""


def ev(d):
    return (d["eval_note"].replace("第四轮学科评估 ", "第四轮学科评估").replace("个学科获评 A+", "个A+")
            .replace(" 学科评估 A+", "学科评估A+").replace(" 学科评估A+", "学科评估A+"))


def ev_entry(d, entry, tail):
    e = ev(d)
    return e if len(e) + len(entry) > 26 else f"{e}，{entry}{tail}"


def city(d):
    return d["city"].split("·")[0]


def subj(d):
    a, b = d["subject_req"][0]
    if a in ("选科要求", "报考条件"):
        a, b = d["subject_req"][1]
        return f"📚 报考提醒：{b}" if a == "报考条件" else f"📚 报考提醒：{a}{b}"
    if "+" in b:
        return f"📚 选科提醒：{a}要求{b.replace(' ', '')}"
    return f"📚 报考提醒：{a}{b}"


def career(d):
    c = d["career"]
    if c.get("rate_value"):
        return f"💼 {c['rate_year']}{c['rate_label']}{c['rate_value']}"
    return f"💼 {c['title']}：{'、'.join(c['directions'][:2])}"


def pick(pool, i, off=0):
    return pool[(i + off) % len(pool)]


def build(sid, i):
    d = json.loads((ROOT / "data" / f"{sid}.json").read_text(encoding="utf-8"))
    hook, cold, q, entry, hist = H[sid]
    cat = CAT.get(sid, "工")
    name, rk, idt, m = d["name"], rank(d), ident(d), motto(d)
    mil = cat == "军"
    idrk = f"{idt}，{rk}" if rk else idt
    mline = f"📖 {motto_label(d)}：{m}"
    note = motto_note(d)

    # 我的B：冷知识开头式
    b = [cold, f"🎓 {name}，{d['founded']}年建校，{city(d)}", f"🏆 {ev(d)}" + (f"，{rk}" if rk else ""),
         f"📜 「{m}」" + (f"，{note}" if note else ""), "✨ " + pick(ENC_B, i)]
    # 我的C：家长干货式
    cta = CTA_MIL if mil else pick(CTA, i).format(s=d["short"])
    fit = d["fit"][0] + (f"，{d['fit'][1]}" if len(d["fit"][0]) + len(d["fit"][1]) <= 16 else "")
    c = [f"🏛️ {name}｜{idrk}", subj(d), f"🎯 适合的孩子：{fit}", career(d), "👉 " + cta]
    # 你的A：标准介绍型
    ua = [TITLE.format(name), f"🏛️ 坐落{city(d)}，{idt}高校，{d['slogan']}",
          "⚙️ " + ev_entry(d, entry, "值得重点了解"), mline, "🌱 " + pick(ENC_A, i)]
    # 你的B：兴趣引导型
    ub = [q, f"🎓 今天认识{name}，{AUD[cat]}",
          f"🏛️ {idt}，{entry}是认识它的入口", LEARN[cat],
          f"📖 「{m}」，{pick(MOTTO_TAIL, i)}", "🌱 " + pick(ENC_Q, i)]
    # 你的C：校史故事型
    uc = [TITLE.format(name), hist, f"⚙️ 如今是{idt}高校，{ev(d)}",
          f"📖 「{m}」" + (f"，{note}" if note else f"，写进了{motto_label(d)}"), "🌱 " + pick(ENC_C, i)]
    # 融合版
    hook_line = {"c": cold, "q": q, "h": hist}[hook]
    close = {"c": ENC_B, "q": ENC_Q, "h": ENC_C}[hook]
    fused = [hook_line, TITLE.format(name), f"🏛️ {idrk}", "⚙️ " + ev_entry(d, entry, "值得关注"), mline,
             "🌱 " + pick(close, i, 3)]
    fused_wed = [TITLE.format(name), f"🏛️ {idrk}", subj(d), f"🎯 适合的孩子：{fit}", career(d), "👉 " + cta]
    j = lambda xs: "\n".join(xs)  # noqa: E731
    return d, {"融合版": j(fused), "融合版·家长干货日": j(fused_wed), "冷知识开头式": j(b),
               "家长干货式": j(c), "标准介绍型": j(ua), "兴趣引导型": j(ub), "校史故事型": j(uc)}, hook


def n(t):
    return len(t.replace("\n", ""))


HOOK_NAME = {"c": "冷知识钩子", "q": "兴趣提问钩子", "h": "校史钩子"}
VERS = ["融合版", "冷知识开头式", "家长干货式", "标准介绍型", "兴趣引导型", "校史故事型"]
SRC = {"融合版": "融合", "冷知识开头式": "Claude方案B", "家长干货式": "Claude方案C",
       "标准介绍型": "你的方案A", "兴趣引导型": "你的方案B", "校史故事型": "你的方案C"}

_ml = openpyxl.load_workbook(ROOT / "data" / "985_211院校总名单.xlsx").active
LOC = {r[1]: (r[3], r[4]) for r in _ml.iter_rows(min_row=2, values_only=True)}  # 校名 -> (省份, 城市)，主校区所在地

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "发布排期（融合版）"
ws.append(["序号", "日期", "星期", "校名", "省份", "城市", "当天版本", "文案（直接复制）", "字数"])
ws2 = wb.create_sheet("六版全集")
ws2.append(["序号", "日期", "星期", "校名", "省份", "城市"] + [f"{v}\n（{SRC[v]}）" for v in VERS])
long_ = []
for i, sid in enumerate(ORDER, 1):
    day = START + dt.timedelta(days=i - 1)
    wk = WEEK[day.weekday()]
    d, t, hook = build(sid, i)
    wed = day.weekday() == 2
    today = t["融合版·家长干货日"] if wed else t["融合版"]
    ws.append([i, day.strftime("%m-%d"), f"周{wk}", d["name"], *LOC[d["name"]],
               "家长干货日" if wed else f"融合版·{HOOK_NAME[hook]}", today, n(today)])
    row = [i, day.strftime("%m-%d"), f"周{wk}", d["name"], *LOC[d["name"]]]
    for v in VERS:
        txt = t["融合版·家长干货日"] if (v == "融合版" and wed) else t[v]
        row.append(txt)
        if n(txt) > 150:
            long_.append((d["name"], v, n(txt)))
    ws2.append(row)

wrap = Alignment(wrap_text=True, vertical="top")
hdr = PatternFill("solid", fgColor="1F3A5F")
for sh, widths in ((ws, [6, 8, 6, 16, 8, 8, 16, 60, 6]), (ws2, [6, 8, 6, 16, 8, 8] + [46] * 6)):
    for c in sh[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for k, w in enumerate(widths):
        sh.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = w
    for r in sh.iter_rows(min_row=2):
        for c in r:
            c.alignment = wrap
        if r[2].value == "周三":
            for c in r[:7]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    sh.freeze_panes = "G2"

ws3 = wb.create_sheet("使用说明")
for line in [
    ["发布节奏", "2026-11-02（周一）起每天 1 所，共 115 天；周三（黄底）为家长干货日，融合版自动换成选科/适合的孩子/去向 + 私信引导"],
    ["融合版结构", "开头钩子（冷知识/兴趣提问/校史，按学校素材选）→ 每天认识一所大学｜校名 → 身份+排名 → 优势学科 → 校训 → 鼓励"],
    ["六版来源", "融合版；Claude方案B 冷知识开头式；Claude方案C 家长干货式；你的方案A 标准介绍型；你的方案B 兴趣引导型；你的方案C 校史故事型"],
    ["省份/城市", "主校区所在地，取自院校总名单（与海报一致）"],
    ["数据口径", "排名、学科、校训、选科、去向等全部取自海报数据（与海报一致）；海军军医、空军军医为「寄语」"],
    ["字数", "不含换行、含表情；建议 ≤150 字，避免朋友圈折叠成「全文」"],
    ["换版本", "想某天换别的版本：到「六版全集」同一行复制对应列即可"],
]:
    ws3.append(line)
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 100
out = ROOT / "data" / "朋友圈文案.xlsx"
wb.save(out)
print(out)
print("超150字：", long_ or "无")
