"""生成实景视频分批计划表：按海报发布顺序，每批 10 所，附每校素材体检结果与特殊处理标记。

用法：python3 scripts/build_batch_plan.py --videos /home/user/china-universities/videos
"""
import argparse
import csv
import collections
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
ORDER = ROOT.parent / "college-posters" / "data" / "发布顺序.xlsx"
DONE = {"浙江大学": "已出实景样片"}
PILOT = ["北京理工大学", "清华大学", "南京大学", "中国科学技术大学"]   # 动效版换实景 / 横屏为主 / 竖屏带原片字幕 / 素材不足 60 秒
PER_BATCH = 10


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--videos", required=True)
    args = ap.parse_args()
    rows = list(csv.DictReader(open(Path(args.videos) / "素材总台账.csv", encoding="utf-8-sig")))
    clips = collections.defaultdict(list)
    for r in rows:
        w, h = map(int, r["分辨率"].split("x"))
        clips[r["学校"]].append({"file": r["文件名"], "dur": float(r["时长秒"]), "w": w, "h": h,
                                 "official": r["来源类型"].startswith("高校官方")})

    order = [r for r in openpyxl.load_workbook(ORDER).active.iter_rows(min_row=2, values_only=True) if r[0]]
    queue = [r for r in order if r[1] not in DONE and r[1] not in PILOT]
    batch_of = {n: "试点批" for n in PILOT}
    for i, r in enumerate(queue):
        batch_of[r[1]] = f"第{i // PER_BATCH + 1}批"
    batch_of.update({n: "已完成" for n in DONE})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "分批计划"
    head = ["发布序号", "校名", "层次", "视频批次", "素材条数", "横屏", "竖屏", "最高清晰度(短边)", "素材总时长(秒)",
            "官方来源", "非官方来源", "处理方式", "需特别注意", "状态"]
    ws.append(head)
    stat = collections.Counter()
    for seq, name, level, *_ in order:
        c = clips.get(name, [])
        land = sum(x["w"] > x["h"] for x in c)
        port = sum(x["h"] > x["w"] for x in c)
        best = max((min(x["w"], x["h"]) for x in c), default=0)
        total = round(sum(x["dur"] for x in c), 1)
        off = sum(x["official"] for x in c)
        mode = "竖屏铺满" if port and not land else ("横屏卡片" if land and not port else "横竖混剪")
        notes = []
        if best < 720:
            notes.append("清晰度偏低：多用卡片和虚化背景")
        if total < 60:
            notes.append("素材不足60秒：信息段用动效补")
        if off < len(c):
            notes.append(f"含{len(c) - off}条非官方素材：片尾署名")
        stat[batch_of[name]] += 1
        ws.append([seq, name, level, batch_of[name], len(c), land, port, best, total, off, len(c) - off, mode,
                   "；".join(notes), DONE.get(name, "待制作")])

    fills = {"试点批": "FFF2CC", "已完成": "E3F1E8"}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0A2A5C")
    for row in ws.iter_rows(min_row=2):
        f = fills.get(row[3].value)
        if f:
            for c in row:
                c.fill = PatternFill("solid", fgColor=f)
        row[12].alignment = Alignment(wrap_text=True)
    for col, w in zip("ABCDEFGHIJKLMN", [8, 20, 9, 9, 8, 6, 6, 14, 13, 8, 9, 10, 40, 22]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("批次汇总")
    ws2.append(["批次", "学校数", "学校"])
    for b in ["已完成", "试点批"] + [f"第{i}批" for i in range(1, 13)]:
        names = [r[1] for r in order if batch_of[r[1]] == b]
        if names:
            ws2.append([b, len(names), "、".join(names)])
    ws2.column_dimensions["C"].width = 120
    for c in ws2[1]:
        c.font = Font(bold=True)
    for row in ws2.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True)

    dst = ROOT / "data" / "实景视频分批计划.xlsx"
    wb.save(dst)
    print(dst, dict(stat))


if __name__ == "__main__":
    main()
